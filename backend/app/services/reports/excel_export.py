"""Excel çıktıları: takip Excel'ine benzer çizelge, borçlu listesi, daire ekstresi."""

from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_TR_MONTHS = ["OCAK", "ŞUBAT", "MART", "NİSAN", "MAYIS", "HAZİRAN", "TEMMUZ", "AĞUSTOS", "EYLÜL", "EKİM", "KASIM", "ARALIK"]
_FILLS = {
    "paid": PatternFill("solid", fgColor="D9F2E3"),
    "partial": PatternFill("solid", fgColor="FFF1CC"),
    "unpaid": PatternFill("solid", fgColor="F9D6D5"),
}
_HEAD = Font(bold=True)
_THIN = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
_MONEY = '#,##0.00'


def _label(period: str) -> str:
    y, m = period.split("-")
    return f"{_TR_MONTHS[int(m) - 1]} {y}"


def _num(v: str | None) -> float | None:
    if v in (None, ""):
        return None
    return float(Decimal(v))


def grid_workbook(site_name: str, grid: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "TAHSİLAT"
    periods = grid["periods"]
    ws.cell(row=1, column=1, value=f"{site_name} — Aidat tahsilat çizelgesi").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="Hücre: o ay gönderilen ödemeler. Renk: yeşil ödendi, sarı eksik, kırmızı ödenmedi (o ayın tahakkukuna göre).")
    r = 4
    for block in grid["blocks"]:
        ws.cell(row=r, column=1, value="No").font = _HEAD
        ws.cell(row=r, column=2, value=f"{block['name']} BLOK").font = _HEAD
        for j, p in enumerate(periods):
            c = ws.cell(row=r, column=3 + j, value=_label(p))
            c.font = _HEAD
            c.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3 + len(periods), value="BAKİYE").font = _HEAD
        ws.cell(row=r, column=4 + len(periods), value="GECİKEN AY").font = _HEAD
        r += 1
        for u in block["units"]:
            ws.cell(row=r, column=1, value=u["number"])
            ws.cell(row=r, column=2, value=" / ".join(u["names"]) or u["code"])
            for j, p in enumerate(periods):
                cell = u["cells"].get(p)
                c = ws.cell(row=r, column=3 + j)
                c.border = _THIN
                if not cell:
                    continue
                paid = _num(cell["paid"])
                c.value = paid if paid else None
                c.number_format = _MONEY
                if cell["status"] in _FILLS:
                    c.fill = _FILLS[cell["status"]]
            bc = ws.cell(row=r, column=3 + len(periods), value=_num(u["balance"]))
            bc.number_format = _MONEY
            ws.cell(row=r, column=4 + len(periods), value=u["overdue_months"])
            r += 1
        r += 1
    ws.cell(row=r, column=2, value="TOPLAM").font = _HEAD
    for j, p in enumerate(periods):
        c = ws.cell(row=r, column=3 + j, value=_num(grid["totals"].get(p)))
        c.number_format = _MONEY
        c.font = _HEAD
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 38
    for j in range(len(periods) + 2):
        ws.column_dimensions[get_column_letter(3 + j)].width = 13
    ws.freeze_panes = "C5"
    return _save(wb)


def debtors_workbook(site_name: str, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BORÇLULAR"
    headers = ["Daire", "Sakin(ler)", "Tahakkuk", "Ödenen", "Bakiye", "Geciken ay", "En eski ödenmemiş", "Grup"]
    ws.append([f"{site_name} — Borçlu listesi"])
    ws.append(headers)
    for c in ws[2]:
        c.font = _HEAD
    for row in rows:
        ws.append([row["code"], " / ".join(row["names"]), _num(row["charged"]), _num(row["paid"]), _num(row["balance"]),
                   row["overdue_months"], row["oldest_unpaid_period"], row["bucket"]])
    for col in ("C", "D", "E"):
        for c in ws[col][2:]:
            c.number_format = _MONEY
    for i, w in enumerate([8, 40, 14, 14, 14, 11, 18, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return _save(wb)


def statement_workbook(site_name: str, unit_code: str, entries: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "EKSTRE"
    ws.append([f"{site_name} — {unit_code} hesap ekstresi"])
    ws.append(["Tarih", "Dönem", "Tür", "Açıklama", "Borç", "Alacak", "Bakiye"])
    for c in ws[2]:
        c.font = _HEAD
    for e in entries:
        ws.append([e["date"], e["period"], e["type"], e["description"], float(e["debit"]) or None,
                   float(e["credit"]) or None, float(e["balance"])])
    for col in ("E", "F", "G"):
        for c in ws[col][2:]:
            c.number_format = _MONEY
    for i, w in enumerate([12, 9, 10, 70, 13, 13, 13], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
